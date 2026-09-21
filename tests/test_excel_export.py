import json
from pathlib import Path
import re
import openpyxl
import pytest

from src.excel_exporter import build_active_corpus_excel_report
from src.ragas_eval import generate_synthetic_corpus_queries

KNOWN_MANUSCRIPT_STEMS = [
    "2601.05264",
    "2604.01733",
    "2605.28222",
    "2606.24937",
    "2606.29090",
]

CHALLENGE_TIERS = {
    "Empirical Fact",
    "Methodological Synthesis",
    "Cross-Document Disambiguation",
}

def test_build_active_corpus_excel_report(tmp_path):
    active_cache_file = Path("data/cache_active_corpus_benchmark.json")
    if not active_cache_file.exists():
        pytest.skip("No active corpus cache found for test")
    
    data = json.loads(active_cache_file.read_text(encoding="utf-8"))
    out_file = tmp_path / "test_active_eval.xlsx"
    
    result_path = build_active_corpus_excel_report(data, out_file, indexed_chunks_count=5376)
    assert result_path.exists()
    assert result_path.stat().st_size > 5000
    
    wb = openpyxl.load_workbook(str(result_path))
    expected_sheets = [
        "Executive Summary",
        "Query-Level Evaluations",
        "Category Breakdown",
        "Ingested Manuscripts Manifest"
    ]
    assert wb.sheetnames == expected_sheets
    
    ws1 = wb["Executive Summary"]
    assert "Active Ingested Corpus Benchmark Report" in str(ws1["B2"].value)
    
    ws2 = wb["Query-Level Evaluations"]
    assert ws2.max_row >= 10
    
    ws3 = wb["Category Breakdown"]
    assert ws3.max_row >= 4
    
    ws4 = wb["Ingested Manuscripts Manifest"]
    assert ws4.max_row >= 5


def test_synthetic_query_generator_zero_leakage():
    """Verify synthetic question generation produces natural queries without leaking document filenames."""
    sample_chunks = [
        {
            "chunk_id": "c1",
            "filename": "2601.05264v1.pdf",
            "page_number": 1,
            "text": "The TruLens framework evaluates groundedness and hallucination suppression across multi-hop reasoning chains.",
        },
        {
            "chunk_id": "c2",
            "filename": "2604.01733v1.pdf",
            "page_number": 5,
            "text": "Table 3 compares BM25 vs dense embeddings on FinQA benchmark yielding 83.4 percent recall with cross-encoders.",
        },
        {
            "chunk_id": "c3",
            "filename": "2606.24937v2.pdf",
            "page_number": 12,
            "text": "Context compression algorithms truncate past dialogue turns when the token budget exceeds the prompt window.",
        },
    ]
    queries = generate_synthetic_corpus_queries(sample_chunks, num_queries=6, llm_instance=None)
    assert len(queries) == 6
    
    for item in queries:
        q_text = item["query"]
        # Filename leakage assertion
        assert ".pdf" not in q_text.lower(), f"Leaked file extension in query: {q_text}"
        for stem in KNOWN_MANUSCRIPT_STEMS:
            assert stem not in q_text, f"Leaked manuscript stem '{stem}' in query: {q_text}"
        # Tier validation
        assert item.get("category") in CHALLENGE_TIERS, f"Invalid challenge tier: {item.get('category')}"
        # Natural syntax assertion
        assert q_text.endswith("?"), f"Query should end with question mark: {q_text}"
        assert len(q_text.split()) >= 5, f"Query too brief: {q_text}"


def test_no_filename_leakage_in_queries(tmp_path):
    """Verify that no query in the active benchmark workbook or cache leaks PDF filenames."""
    active_cache_file = Path("data/cache_active_corpus_benchmark.json")
    if not active_cache_file.exists():
        pytest.skip("No active corpus cache found for test")
        
    data = json.loads(active_cache_file.read_text(encoding="utf-8"))
    out_file = tmp_path / "test_no_leakage.xlsx"
    build_active_corpus_excel_report(data, out_file, indexed_chunks_count=5376)
    
    wb = openpyxl.load_workbook(str(out_file))
    ws2 = wb["Query-Level Evaluations"]
    
    # Sheet 2: Header is row 5, col 5 is "Evaluated Query Text"
    header_col5 = ws2.cell(row=5, column=5).value
    assert header_col5 == "Evaluated Query Text", f"Expected col 5 to be 'Evaluated Query Text', got '{header_col5}'"
    
    query_count = 0
    for r in range(6, ws2.max_row + 1):
        q_text = str(ws2.cell(row=r, column=5).value or "")
        if not q_text or q_text == "None":
            continue
        query_count += 1
        assert ".pdf" not in q_text.lower(), f"Row {r} leaked '.pdf': {q_text}"
        for stem in KNOWN_MANUSCRIPT_STEMS:
            assert stem not in q_text, f"Row {r} leaked manuscript stem '{stem}': {q_text}"
            
    assert query_count >= 10, f"Expected at least 10 query rows, found {query_count}"


def test_challenge_tier_distribution(tmp_path):
    """Verify queries span the 3 challenge tiers and are mapped into Sheet 2 and Sheet 3."""
    active_cache_file = Path("data/cache_active_corpus_benchmark.json")
    if not active_cache_file.exists():
        pytest.skip("No active corpus cache found for test")
        
    data = json.loads(active_cache_file.read_text(encoding="utf-8"))
    out_file = tmp_path / "test_tiers.xlsx"
    build_active_corpus_excel_report(data, out_file, indexed_chunks_count=5376)
    
    wb = openpyxl.load_workbook(str(out_file))
    ws2 = wb["Query-Level Evaluations"]
    
    # Sheet 2: Header is row 5, col 3 is "Challenge Tier"
    header_col3 = ws2.cell(row=5, column=3).value
    assert header_col3 == "Challenge Tier", f"Expected col 3 to be 'Challenge Tier', got '{header_col3}'"
    
    found_tiers = set()
    for r in range(6, ws2.max_row + 1):
        tier_val = str(ws2.cell(row=r, column=3).value or "")
        if tier_val and tier_val != "None":
            found_tiers.add(tier_val)
            
    assert found_tiers.issubset(CHALLENGE_TIERS), f"Unexpected tiers: {found_tiers - CHALLENGE_TIERS}"
    assert len(found_tiers) >= 2, f"Expected multi-tier representation, found: {found_tiers}"
    
    # Sheet 3: Category Breakdown
    ws3 = wb["Category Breakdown"]
    breakdown_rows = []
    for r in range(6, ws3.max_row + 1):
        cat_name = str(ws3.cell(row=r, column=2).value or "")
        q_count = ws3.cell(row=r, column=3).value
        if cat_name and cat_name != "None":
            breakdown_rows.append((cat_name, q_count))
            
    assert len(breakdown_rows) >= 3, f"Expected at least 3 breakdown categories, found {len(breakdown_rows)}"
    for cat_name, q_count in breakdown_rows:
        assert q_count is not None and int(q_count) > 0, f"Category '{cat_name}' has non-positive count: {q_count}"


def test_mathematical_consistency_summary_vs_queries(tmp_path):
    """Verify micro-averaged top-5 hit rate in Sheet 2 matches Pareto Scorecard in Sheet 1."""
    active_cache_file = Path("data/cache_active_corpus_benchmark.json")
    if not active_cache_file.exists():
        pytest.skip("No active corpus cache found for test")
        
    data = json.loads(active_cache_file.read_text(encoding="utf-8"))
    out_file = tmp_path / "test_math_consistency.xlsx"
    build_active_corpus_excel_report(data, out_file, indexed_chunks_count=5376)
    
    wb = openpyxl.load_workbook(str(out_file))
    ws1 = wb["Executive Summary"]
    ws2 = wb["Query-Level Evaluations"]
    
    # Read Sheet 1 reported Recall@5
    # Row 7: Baseline (Hybrid) -> col 4
    # Row 8: Shallow L-2 Reranker -> col 4
    # Row 9: Deep L-6 Reranker -> col 4
    hy_rec_str = str(ws1.cell(row=7, column=4).value or "0.0%")
    l2_rec_str = str(ws1.cell(row=8, column=4).value or "0.0%")
    l6_rec_str = str(ws1.cell(row=9, column=4).value or "0.0%")
    
    hy_recall_reported = float(hy_rec_str.replace("%", "")) / 100.0
    l2_recall_reported = float(l2_rec_str.replace("%", "")) / 100.0
    l6_recall_reported = float(l6_rec_str.replace("%", "")) / 100.0
    
    # Compute from Sheet 2 query rows
    # Col 6: Hybrid Rank, Col 7: L-2 Rerank, Col 8: L-6 Rerank
    total_queries = 0
    hy_hits = 0
    l2_hits = 0
    l6_hits = 0
    
    def parse_rank_is_top5(val):
        s = str(val or "").strip().lstrip("#")
        if s.isdigit():
            return 1 <= int(s) <= 5
        return False
        
    for r in range(6, ws2.max_row + 1):
        q_id = ws2.cell(row=r, column=2).value
        if not q_id:
            continue
        total_queries += 1
        
        h_rank = ws2.cell(row=r, column=6).value
        l2_rank = ws2.cell(row=r, column=7).value
        l6_rank = ws2.cell(row=r, column=8).value
        
        if parse_rank_is_top5(h_rank):
            hy_hits += 1
        if parse_rank_is_top5(l2_rank):
            l2_hits += 1
        if parse_rank_is_top5(l6_rank):
            l6_hits += 1
            
    assert total_queries > 0, "No query rows found in Sheet 2"
    
    micro_hy_rec = hy_hits / total_queries
    micro_l2_rec = l2_hits / total_queries
    micro_l6_rec = l6_hits / total_queries
    
    # Tolerance of 0.02 allows for rounding to 1 decimal percentage place (e.g. 88.0%)
    assert abs(micro_hy_rec - hy_recall_reported) <= 0.02, (
        f"Hybrid Recall mismatch: Sheet 1 reports {hy_recall_reported:.3f}, Sheet 2 micro-average is {micro_hy_rec:.3f}"
    )
    assert abs(micro_l2_rec - l2_recall_reported) <= 0.02, (
        f"L-2 Recall mismatch: Sheet 1 reports {l2_recall_reported:.3f}, Sheet 2 micro-average is {micro_l2_rec:.3f}"
    )
    assert abs(micro_l6_rec - l6_recall_reported) <= 0.02, (
        f"L-6 Recall mismatch: Sheet 1 reports {l6_recall_reported:.3f}, Sheet 2 micro-average is {micro_l6_rec:.3f}"
    )


def test_significance_p_value_honest_bounds(tmp_path):
    """Verify paired Student's t-test p-values are valid probabilities strictly in [0.0, 1.0]."""
    active_cache_file = Path("data/cache_active_corpus_benchmark.json")
    if not active_cache_file.exists():
        pytest.skip("No active corpus cache found for test")
        
    data = json.loads(active_cache_file.read_text(encoding="utf-8"))
    out_file = tmp_path / "test_significance.xlsx"
    build_active_corpus_excel_report(data, out_file, indexed_chunks_count=5376)
    
    wb = openpyxl.load_workbook(str(out_file))
    ws1 = wb["Executive Summary"]
    
    # Locate Section 3 Statistical Hypothesis Testing in Sheet 1
    t_stat_val = None
    p_val_val = None
    
    for r in range(15, ws1.max_row + 1):
        label = str(ws1.cell(row=r, column=2).value or "")
        val = str(ws1.cell(row=r, column=3).value or "")
        if "Test Statistic" in label:
            t_stat_val = float(val)
        elif "p-value" in label:
            # e.g., "p = 0.0013 < 0.05" or "0.0013"
            m = re.search(r"p\s*=\s*([0-9.]+)", val)
            if m:
                p_val_val = float(m.group(1))
            else:
                try:
                    p_val_val = float(val)
                except ValueError:
                    pass
                    
    assert t_stat_val is not None, "Test statistic (t) not found in Sheet 1"
    assert p_val_val is not None, "p-value not found in Sheet 1"
    
    # Rigorous probabilistic validity
    assert 0.0 <= p_val_val <= 1.0, f"p-value out of bounds: {p_val_val}"
    assert abs(t_stat_val) < 100.0, f"Absurd t-statistic value: {t_stat_val}"
