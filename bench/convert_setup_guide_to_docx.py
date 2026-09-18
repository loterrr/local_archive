import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DOCX = ROOT / "SETUP_GUIDE.docx"

doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

COLOR_PRIMARY = RGBColor(31, 78, 121)    # Deep Navy
COLOR_SECONDARY = RGBColor(46, 117, 182) # Blue
COLOR_DARK = RGBColor(30, 41, 59)        # Slate Dark
COLOR_TEXT = RGBColor(51, 51, 51)        # Body Text
COLOR_MUTED = RGBColor(100, 116, 139)    # Gray

normal_style = doc.styles['Normal']
normal_style.font.name = 'Calibri'
normal_style.font.size = Pt(11)
normal_style.font.color.rgb = COLOR_TEXT
normal_style.paragraph_format.line_spacing = 1.15
normal_style.paragraph_format.space_after = Pt(6)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''<w:tcMar {nsdecls("w")}>
        <w:top w:w="{top}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>''')
    tcPr.append(tcMar)

def add_header(text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.bold = True
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(15)
        run.font.color.rgb = COLOR_PRIMARY
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="1F4E79"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(12.5)
        run.font.color.rgb = COLOR_SECONDARY
    return p

def add_code_block(code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(15, 23, 42)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)

# Title
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(2)
run = p.add_run("LocalRAG: New Device Deployment & Setup Guide")
run.bold = True
run.font.size = Pt(22)
run.font.color.rgb = COLOR_PRIMARY

sub = doc.add_paragraph("Complete step-by-step instructions for Windows, macOS, and Linux")
sub.runs[0].font.italic = True
sub.runs[0].font.color.rgb = COLOR_MUTED

# 1. Requirements
add_header("1. System Requirements & Prerequisites", level=1)
doc.add_paragraph("• Operating System: Windows 10/11, macOS 12+, or Ubuntu 20.04+\n• RAM: 8 GB minimum, 16 GB recommended for smooth LLM generation\n• CPU: 4+ cores (100% CPU-optimized, no dedicated GPU required)\n• Disk: 10 GB free space for models and indexes")

doc.add_paragraph("Required Software to Install:")
doc.add_paragraph("1. Python 3.10 or 3.11 (Check 'Add python.exe to PATH' during Windows setup)")
doc.add_paragraph("2. Ollama from https://ollama.com (For local offline LLM generation)")
doc.add_paragraph("3. Optional: Tesseract OCR (Only needed for scanning image-only PDFs)")

# 2. Setup
add_header("2. Quick Setup Steps", level=1)

add_header("Step 1: Copy Project Files", level=2)
doc.add_paragraph("Copy the entire localrag project folder to your new computer. Keep data/indexes/current/ so you don't need to rebuild the index. Do NOT copy .venv/.")

add_header("Step 2: Windows 1-Click Setup", level=2)
doc.add_paragraph("On Windows, you can simply double-click setup_new_device.bat in the root folder, or run in PowerShell:")
add_code_block("Set-ExecutionPolicy -Scope Process -ExecutionPolicy Bypass\npython -m venv .venv\n.venv\\Scripts\\activate\npython -m pip install --upgrade pip\npip install -r requirements.txt\ncopy .env.example .env")

add_header("Step 3: macOS / Linux Setup", level=2)
doc.add_paragraph("On macOS or Linux, run:")
add_code_block("python3 -m venv .venv\nsource .venv/bin/activate\npython -m pip install --upgrade pip\npip install -r requirements.txt\ncp .env.example .env")

add_header("Step 4: Ollama Model Setup", level=2)
doc.add_paragraph("Start Ollama and pull the lightweight 3.8B model:")
add_code_block("ollama pull phi3.5:latest\nollama run phi3.5:latest \"Hello!\"")

add_header("Step 5: Launch the Application", level=2)
doc.add_paragraph("Start the server:")
add_code_block("python app/server.py")
doc.add_paragraph("Then open your browser at: http://localhost:3000")

# 3. Verification
add_header("3. Verification & Benchmark Testing", level=1)
doc.add_paragraph("1. Check UI Status: Ollama: Connected (phi3.5:latest), Chunks: 4,995 (20 manuscripts).\n2. Run Unit Tests: Run 'pytest tests/' in terminal (6 tests pass in ~1.2s).\n3. Live Benchmark: In the web UI, go to 'Empirical Benchmark' and click 'Run Benchmark'. In ~20 seconds, it evaluates 50 queries live across FAISS, BM25, and Cross-Encoder.")

doc.save(str(OUTPUT_DOCX))
print(f"[DONE] Word Document successfully created at: {OUTPUT_DOCX}")
