import json
import os
import sys
import time
from pathlib import Path
from typing import Sequence

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from src.config import Settings
from src.evaluation import recall_at_k, precision_at_k, reciprocal_rank_at_k
from src.retrieval import HybridIndex
from src.reranker import CrossEncoderReranker

def clean_cell_text(val) -> str:
    if val is None:
        return ""
    s = str(val)
    return ILLEGAL_CHARACTERS_RE.sub("", s)

def evaluate_rank_list(rankings: dict[str, Sequence[str]], queries: list[dict], k: int = 5) -> dict:
    rows = []
    for q in queries:
        qid = q["query_id"]
        rel = set(q["relevant_chunk_ids"])
        ranked = list(rankings.get(qid, []))
        first = next((i for i, cid in enumerate(ranked[:k], 1) if cid in rel), None)
        rows.append({
            "query_id": qid,
            "category": q.get("category", "General"),
            "query": q["query"],
            "first_relevant_rank": first,
            "recall_at_k": recall_at_k(ranked, rel, k),
            "mrr_at_k": reciprocal_rank_at_k(ranked, rel, k),
            "precision_at_k": precision_at_k(ranked, rel, k),
        })
    n = len(rows)
    return {
        "k": k,
        "num_queries": n,
        "recall_at_k": round(sum(r["recall_at_k"] for r in rows) / n, 4) if n else 0.0,
        "mrr_at_k": round(sum(r["mrr_at_k"] for r in rows) / n, 4) if n else 0.0,
        "precision_at_k": round(sum(r["precision_at_k"] for r in rows) / n, 4) if n else 0.0,
        "details": rows,
    }

def run_evaluation():
    settings = Settings()
    index_dir = ROOT / "data" / "indexes" / "current"
    dataset_file = ROOT / "data" / "evaluation_deep_research_50.json"

    print("Loading index & queries...")
    index = HybridIndex.load(index_dir, settings.embedding_model)
    queries = json.loads(dataset_file.read_text(encoding="utf-8"))["queries"]
    reranker = CrossEncoderReranker(settings.reranker_model)

    hybrid_rankings = {}
    reranked_rankings = {}
    query_timings = []
    k_candidates = 50
    k_rerank_max = 20

    cache_file = ROOT / "data" / "cache_deep_eval.json"
    if cache_file.exists():
        print(f"Loading pre-computed rankings from {cache_file}...")
        cached = json.loads(cache_file.read_text(encoding="utf-8"))
        hybrid_rankings = cached["hybrid_rankings"]
        reranked_rankings = cached["reranked_rankings"]
        query_timings = cached["query_timings"]
        reranker_metrics = cached["reranker_metrics"]
    else:
        print(f"Executing search and reranking on {len(queries)} queries (candidate_k={k_candidates}, top_k={k_rerank_max})...")
        t_start = time.perf_counter()
        for idx, q in enumerate(queries, 1):
            t0 = time.perf_counter()
            candidates = index.search(
                q["query"],
                dense_k=settings.dense_k,
                final_k=k_rerank_max,
                rrf_k=settings.rrf_k,
                candidate_k=k_candidates
            )
            hybrid_rankings[q["query_id"]] = [c.chunk_id for c, *_ in candidates]
            rr = reranker.rerank(q["query"], candidates, top_k=k_rerank_max, batch_size=settings.reranker_batch_size)
            reranked_rankings[q["query_id"]] = [x.chunk.chunk_id for x in rr]

            elapsed = (time.perf_counter() - t0) * 1000
            query_timings.append(elapsed)
            if idx % 10 == 0 or idx == len(queries):
                print(f"  Processed {idx}/{len(queries)} queries ({elapsed:.1f} ms for last query)")

        reranker_metrics = reranker.metrics()
        cache_file.write_text(json.dumps({
            "hybrid_rankings": hybrid_rankings,
            "reranked_rankings": reranked_rankings,
            "query_timings": query_timings,
            "reranker_metrics": reranker_metrics
        }, indent=2), encoding="utf-8")

        total_time_s = time.perf_counter() - t_start
        print(f"Retrieval and reranking completed in {total_time_s:.2f} seconds.")

    # Compute Multi-K metrics: K = 1, 5, 10, 20
    k_values = [1, 5, 10, 20]
    multi_k_metrics = {}
    for k in k_values:
        hy_k = evaluate_rank_list(hybrid_rankings, queries, k=k)
        rr_k = evaluate_rank_list(reranked_rankings, queries, k=k)
        multi_k_metrics[k] = {"hybrid": hy_k, "reranked": rr_k}

    # Compute Category metrics at K = 10 (Deep Research standard)
    categories = ["Hard Distractor", "Empirical Fact", "Synthesis"]
    category_metrics = {}
    for cat in categories:
        cat_queries = [q for q in queries if q.get("category") == cat]
        hy_cat = evaluate_rank_list(hybrid_rankings, cat_queries, k=10)
        rr_cat = evaluate_rank_list(reranked_rankings, cat_queries, k=10)
        category_metrics[cat] = {
            "num_queries": len(cat_queries),
            "hybrid": hy_cat,
            "reranked": rr_cat
        }

    # Build Excel report
    out_excel = ROOT / "data" / "retrieval_evaluation_comparison.xlsx"
    build_excel_report(queries, multi_k_metrics, category_metrics, hybrid_rankings, reranked_rankings, query_timings, reranker_metrics, settings, out_excel, k_candidates=k_candidates)

    # Also update data/evaluation_labeled.json with this 50-query dataset
    labeled_json = ROOT / "data" / "evaluation_labeled.json"
    labeled_json.write_text(json.dumps({"queries": queries}, indent=2), encoding="utf-8")
    print(f"Updated {labeled_json} with 50 deep research queries.")

def build_excel_report(queries, multi_k_metrics, category_metrics, hybrid_rankings, reranked_rankings, query_timings, rerank_metrics, settings, out_path: Path, k_candidates: int = 50):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    font_family = "Segoe UI"
    title_font = Font(name=font_family, size=15, bold=True, color="1F4E79")
    subtitle_font = Font(name=font_family, size=10, italic=True, color="595959")
    section_font = Font(name=font_family, size=12, bold=True, color="1F4E79")
    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    cell_font = Font(name=font_family, size=9.5, color="000000")
    bold_cell_font = Font(name=font_family, size=9.5, bold=True, color="000000")

    navy_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    blue_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    slate_fill = PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    improved_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # soft green
    recovered_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # soft purple/blue
    unchanged_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    regression_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # soft amber

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # -------------------------------------------------------------
    # SHEET 1: Executive Summary
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Executive Summary")
    ws1.views.sheetView[0].showGridLines = True

    ws1["B2"] = "Deep Research Retrieval & Reranking Benchmark (50 Queries)"
    ws1["B2"].font = title_font
    ws1["B3"] = f"Comparative evaluation with distractor traps across 5 academic RAG papers | Evaluated: {time.strftime('%Y-%m-%d %H:%M')}"
    ws1["B3"].font = subtitle_font

    ws1["B5"] = "1. Multi-Cutoff Retrieval Performance (K = 1, 5, 10, 20 | Candidate Pool = 50)"
    ws1["B5"].font = section_font

    headers_s1 = ["Cutoff (K)", "Metric", "Hybrid (Dense+BM25)", "Hybrid + Cross-Encoder", "Absolute Delta", "Relative Change", "Research Impact"]
    for col_idx, h in enumerate(headers_s1, start=2):
        cell = ws1.cell(row=6, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws1.row_dimensions[6].height = 24

    row_curr = 7
    for k in [1, 5, 10, 20]:
        hy = multi_k_metrics[k]["hybrid"]
        rr = multi_k_metrics[k]["reranked"]

        metrics = [
            ("Recall@K", hy["recall_at_k"], rr["recall_at_k"], True, "Proportion of queries with at least one ground-truth passage in Top-K"),
            ("MRR@K", hy["mrr_at_k"], rr["mrr_at_k"], False, "Reciprocal rank of first ground-truth passage (Primacy in context)"),
            ("Precision@K", hy["precision_at_k"], rr["precision_at_k"], True, "Concentration of relevant passages among Top-K context segments"),
        ]

        for m_name, h_val, r_val, is_pct, note in metrics:
            diff = r_val - h_val
            pct = (diff / h_val) if h_val > 0 else 0.0

            ws1.cell(row=row_curr, column=2, value=f"K = {k}")
            ws1.cell(row=row_curr, column=3, value=m_name)
            
            c_hy = ws1.cell(row=row_curr, column=4, value=h_val)
            c_rr = ws1.cell(row=row_curr, column=5, value=r_val)
            c_diff = ws1.cell(row=row_curr, column=6, value=diff)
            c_pct = ws1.cell(row=row_curr, column=7, value=pct)
            ws1.cell(row=row_curr, column=8, value=note)

            for c in [c_hy, c_rr, c_diff]:
                if is_pct:
                    c.number_format = "0.0%"
                else:
                    c.number_format = "0.000"
            c_pct.number_format = "+0.0%;-0.0%;0.0%"

            for c_idx in range(2, 9):
                cell = ws1.cell(row=row_curr, column=c_idx)
                cell.font = cell_font
                cell.border = thin_border
                if c_idx in [2, 3]:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_idx in [4, 5, 6, 7]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                if row_curr % 2 == 0:
                    cell.fill = zebra_fill
            row_curr += 1

    # Section 2: Latency & Telemetry
    row_curr += 2
    ws1.cell(row=row_curr, column=2, value="2. Latency & Telemetry Profile").font = section_font
    row_curr += 1

    headers_s2 = ["Pipeline Component", "Implementation / Model", "Avg Latency (ms)", "Total Run Time", "Notes / Sizing"]
    for col_idx, h in enumerate(headers_s2, start=2):
        cell = ws1.cell(row=row_curr, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = blue_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws1.row_dimensions[row_curr].height = 24
    row_curr += 1

    avg_ms = sum(query_timings) / len(query_timings) if query_timings else 0.0
    rerank_avg_ms = rerank_metrics.get("avg_latency_ms", 0.0)
    retrieval_avg_ms = max(0.0, avg_ms - rerank_avg_ms)

    telemetry_rows = [
        ("Hybrid Retrieval", f"{settings.embedding_model} + BM25", retrieval_avg_ms, f"{retrieval_avg_ms * len(queries) / 1000:.1f} s", f"Dense FAISS + BM25Okapi candidate pool ({k_candidates} chunks)"),
        ("Cross-Encoder Reranker", f"{settings.reranker_model}", rerank_avg_ms, f"{rerank_avg_ms * len(queries) / 1000:.1f} s", f"MS-MARCO MiniLM-L6 cross-attention ({rerank_metrics.get('total_pairs', 0)} pairs evaluated)"),
        ("Complete End-to-End", "Hybrid + Reranking Pipeline", avg_ms, f"{sum(query_timings) / 1000:.1f} s", f"{len(queries)} total research queries with distractors"),
    ]

    for stage, comp, lat, tot, note in telemetry_rows:
        ws1.cell(row=row_curr, column=2, value=stage)
        ws1.cell(row=row_curr, column=3, value=comp)
        c_lat = ws1.cell(row=row_curr, column=4, value=lat)
        c_lat.number_format = "#,##0.0 \"ms\""
        ws1.cell(row=row_curr, column=5, value=tot)
        ws1.cell(row=row_curr, column=6, value=note)

        for c_idx in range(2, 7):
            cell = ws1.cell(row=row_curr, column=c_idx)
            cell.font = bold_cell_font if "Complete" in stage else cell_font
            cell.border = thin_border
            if c_idx == 4:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        row_curr += 1

    # -------------------------------------------------------------
    # SHEET 2: Category Breakdown (Distractors vs Facts vs Synthesis)
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Category Breakdown")
    ws2.views.sheetView[0].showGridLines = True

    ws2["B2"] = "Benchmark Performance by Query Category (Cutoff K = 10)"
    ws2["B2"].font = title_font
    ws2["B3"] = "Testing how Cross-Encoder handles Hard Distractor traps vs. Empirical Fact queries vs. Architectural Synthesis queries."
    ws2["B3"].font = subtitle_font

    headers_s2 = [
        "Category", "Query Count", "Description / Nature of Challenge",
        "Hybrid Recall@10", "Reranked Recall@10", "Recall Delta",
        "Hybrid MRR@10", "Reranked MRR@10", "MRR Delta"
    ]
    for col_idx, h in enumerate(headers_s2, start=2):
        cell = ws2.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws2.row_dimensions[5].height = 28

    cat_descriptions = {
        "Hard Distractor": "High lexical overlap across multiple papers designed to mislead BM25 / bi-encoder keyword matchers",
        "Empirical Fact": "Specific benchmark scores, mathematical hyperparameters, tables, and exact experimental settings",
        "Synthesis": "Multi-hop mechanisms, cross-component interactions, and architectural model design choices",
    }

    row_curr = 6
    for cat in ["Hard Distractor", "Empirical Fact", "Synthesis"]:
        c_data = category_metrics[cat]
        hy = c_data["hybrid"]
        rr = c_data["reranked"]

        ws2.cell(row=row_curr, column=2, value=cat)
        ws2.cell(row=row_curr, column=3, value=c_data["num_queries"]).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=row_curr, column=4, value=cat_descriptions.get(cat, ""))

        c_hr = ws2.cell(row=row_curr, column=5, value=hy["recall_at_k"])
        c_rr = ws2.cell(row=row_curr, column=6, value=rr["recall_at_k"])
        c_rd = ws2.cell(row=row_curr, column=7, value=rr["recall_at_k"] - hy["recall_at_k"])

        c_hm = ws2.cell(row=row_curr, column=8, value=hy["mrr_at_k"])
        c_rm = ws2.cell(row=row_curr, column=9, value=rr["mrr_at_k"])
        c_md = ws2.cell(row=row_curr, column=10, value=rr["mrr_at_k"] - hy["mrr_at_k"])

        c_hr.number_format = "0.0%"
        c_rr.number_format = "0.0%"
        c_rd.number_format = "+0.0%;-0.0%;0.0%"
        c_hm.number_format = "0.000"
        c_rm.number_format = "0.000"
        c_md.number_format = "+0.000;-0.000;0.000"

        for col_idx in range(2, 11):
            cell = ws2.cell(row=row_curr, column=col_idx)
            cell.font = cell_font
            cell.border = thin_border
            if col_idx in [5, 6, 7, 8, 9, 10]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in [2, 3]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        row_curr += 1

    # -------------------------------------------------------------
    # SHEET 3: Per-Query Breakdown (All 50 Queries)
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Per-Query Breakdown")
    ws3.views.sheetView[0].showGridLines = True

    ws3["B2"] = "Complete Query-Level Comparison (50 Queries | K = 10)"
    ws3["B2"].font = title_font
    ws3["B3"] = "Ranks, status movements, and evaluation metrics for each research query."
    ws3["B3"].font = subtitle_font

    headers_s3 = [
        "Query ID", "Category", "Question Text", "Hybrid 1st Rank", "Reranked 1st Rank",
        "Rank Delta", "Status", "Hybrid Recall@10", "Reranked Recall@10", "Hybrid MRR@10", "Reranked MRR@10"
    ]
    for col_idx, h in enumerate(headers_s3, start=2):
        cell = ws3.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws3.row_dimensions[5].height = 28

    hy_details = multi_k_metrics[10]["hybrid"]["details"]
    rr_details = multi_k_metrics[10]["reranked"]["details"]

    row_curr = 6
    for i, q in enumerate(queries):
        h_row = hy_details[i]
        r_row = rr_details[i]

        b_rank = h_row["first_relevant_rank"]
        a_rank = r_row["first_relevant_rank"]

        if b_rank is None and a_rank is not None:
            status = "Recovered"
            status_fill = recovered_fill
            rank_delta = f"N/A -> #{a_rank}"
        elif b_rank is not None and a_rank is not None and a_rank < b_rank:
            status = "Improved"
            status_fill = improved_fill
            rank_delta = f"+{b_rank - a_rank} ranks"
        elif b_rank is not None and a_rank is not None and a_rank > b_rank:
            status = "Regressed"
            status_fill = regression_fill
            rank_delta = f"-{a_rank - b_rank} ranks"
        elif b_rank is not None and a_rank is None:
            status = "Lost"
            status_fill = regression_fill
            rank_delta = "Lost from Top-10"
        else:
            status = "Unchanged"
            status_fill = unchanged_fill
            rank_delta = "0"

        ws3.cell(row=row_curr, column=2, value=q["query_id"])
        ws3.cell(row=row_curr, column=3, value=q.get("category", ""))
        ws3.cell(row=row_curr, column=4, value=clean_cell_text(q["query"]))
        ws3.cell(row=row_curr, column=5, value=f"#{b_rank}" if b_rank else "Unranked")
        ws3.cell(row=row_curr, column=6, value=f"#{a_rank}" if a_rank else "Unranked")
        ws3.cell(row=row_curr, column=7, value=rank_delta)
        ws3.cell(row=row_curr, column=8, value=status)

        c_hr = ws3.cell(row=row_curr, column=9, value=h_row["recall_at_k"])
        c_rr = ws3.cell(row=row_curr, column=10, value=r_row["recall_at_k"])
        c_hm = ws3.cell(row=row_curr, column=11, value=h_row["mrr_at_k"])
        c_rm = ws3.cell(row=row_curr, column=12, value=r_row["mrr_at_k"])

        c_hr.number_format = "0.0%"
        c_rr.number_format = "0.0%"
        c_hm.number_format = "0.000"
        c_rm.number_format = "0.000"

        for col_idx in range(2, 13):
            cell = ws3.cell(row=row_curr, column=col_idx)
            cell.font = cell_font
            cell.border = thin_border
            if col_idx in [2, 3, 5, 6, 7, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            if col_idx == 8:
                cell.fill = status_fill

        row_curr += 1

    # -------------------------------------------------------------
    # SHEET 4: Ground Truth Reference
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Ground Truth Reference")
    ws4.views.sheetView[0].showGridLines = True

    ws4["B2"] = "50-Query Benchmark Ground Truth Mapping"
    ws4["B2"].font = title_font
    ws4["B3"] = "Target passage chunk IDs and text excerpts indexed from the 5 academic papers."
    ws4["B3"].font = subtitle_font

    headers_s4 = ["Query ID", "Category", "Question Text", "Target Chunk ID", "Source PDF Document", "Page #", "Ground Truth Excerpt"]
    for col_idx, h in enumerate(headers_s4, start=2):
        cell = ws4.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws4.row_dimensions[5].height = 26

    chunks_path = ROOT / "data" / "indexes" / "current" / "chunks.json"
    chunk_map = {}
    if chunks_path.exists():
        with open(chunks_path, encoding="utf-8") as f:
            for c in json.load(f):
                chunk_map[c["chunk_id"]] = c

    row_curr = 6
    for q in queries:
        for cid in q["relevant_chunk_ids"]:
            cd = chunk_map.get(cid, {})
            fn = cd.get("filename", "N/A")
            pg = cd.get("page_number", "N/A")
            txt = cd.get("text", "")[:280].replace("\n", " ") + "..."

            ws4.cell(row=row_curr, column=2, value=q["query_id"])
            ws4.cell(row=row_curr, column=3, value=q.get("category", ""))
            ws4.cell(row=row_curr, column=4, value=clean_cell_text(q["query"]))
            ws4.cell(row=row_curr, column=5, value=cid)
            ws4.cell(row=row_curr, column=6, value=fn)
            ws4.cell(row=row_curr, column=7, value=f"Page {pg}")
            ws4.cell(row=row_curr, column=8, value=clean_cell_text(txt))

            for col_idx in range(2, 9):
                cell = ws4.cell(row=row_curr, column=col_idx)
                cell.font = cell_font
                cell.border = thin_border
                if col_idx in [2, 3, 5, 7]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                if row_curr % 2 == 0:
                    cell.fill = zebra_fill
            row_curr += 1

    # Auto-adjust column widths
    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            if col_letter == "A":
                ws.column_dimensions[col_letter].width = 4
                continue
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and "\n" not in val_str:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Successfully saved 4-sheet evaluation workbook at: {out_path}")

if __name__ == "__main__":
    run_evaluation()
